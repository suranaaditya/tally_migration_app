"""Orchestrator — loads a RuleSource + COA, emits a MappedDecision per ledger.

Public API:

    from rgi_migration.mapper import Mapper, JsonFileRuleSource, load_coa
    rs  = JsonFileRuleSource(Path("docs/seed_plan.json"))
    coa = load_coa(Path("rgi_migration/tests/fixtures/cacspu_erpnext_coa_real.csv"))
    m   = Mapper(rs, coa, abbr="CACSPU")
    decisions = m.map_all(parsed_tb.ledgers)

Resolution order inside `Mapper.resolve`:

    1. Structural validator — P&L root-type exclusion (validators.py §a)
    2. Anti-pattern match lookup (informational; applied to candidates below)
    3. Positive rule match — exact_ci rules, then pattern-mode rules (tier1_rules.py)
    4. Exact-name fallback against the ERP COA
    5. Anti-pattern-only fallback (rule matched, no positive target resolved)
    6. Unmapped

Post-resolution filters applied to any candidate target:
    - Anti-pattern override: if the candidate equals the anti-pattern's
      `forbidden_erpnext_template`, the candidate is refused and we return
      an anti-pattern-blocked Decision (with creation request if the
      anti-pattern carries creation fields).
    - Group-account refusal (validators.py §b): if the candidate is a group,
      return a "Pending Group Account Resolution" Decision.
    - Missing target with `creates_erpnext_account=1`: emit a "Pending
      Account Creation" Decision with full new-account fields, inheriting
      from a paired anti-pattern if one supplies the creation directive.

CLI (python -m rgi_migration.mapper.mapper --help):
    Loads a COA (CSV or XLSX), parses a Tally XML or Excel fixture, resolves
    every ledger, prints a hit-rate summary as JSON. See `_cli` at the
    bottom of this file.
"""

from __future__ import annotations

import argparse
import csv
import json
import logging
import sys
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from rgi_migration.mapper.rule_source import (
    JsonFileRuleSource,
    Rule,
    RuleSource,
)
from rgi_migration.mapper.supplier_source import Supplier, SupplierSource
from rgi_migration.mapper.tier1_rules import (
    find_matching_anti_pattern,
    find_matching_positive_rule,
    resolve_exact_name,
)
from rgi_migration.mapper.tier1_supplier import (
    is_vendor_party_ledger,
    resolve_supplier,
)
from rgi_migration.mapper.validators import (
    ValidatorOutcome,
    group_account_refusal,
    pnl_root_type_exclusion,
)

LOG = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Data classes — inputs and outputs
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class CoaAccount:
    """One row from the target company's ERPNext Chart of Accounts."""

    name: str
    parent_account: str | None
    root_type: str
    is_group: bool
    company_abbr: str = ""


@dataclass(frozen=True)
class MappedDecision:
    """The mapper's output for a single Tally ledger.

    Mirrors the future `Mapping Decision` DocType child-table shape. Fields
    not populated for a given tier are left as their defaults so downstream
    code can inspect with hasattr-free attribute access.
    """

    # Tally-side input (carried through for downstream review UI)
    tally_name: str
    tally_id: str | None
    tally_root_type: str
    opening_dr: float
    opening_cr: float

    # Decision
    tier: str
    proposed_account: str | None
    review_action: str
    matched_rule: str | None
    confidence: float = 0.0

    # Anti-pattern provenance
    anti_pattern_blocked: bool = False
    anti_pattern_rule: str | None = None
    anti_pattern_message: str | None = None

    # Structural exclusion / group-refusal / target-missing reasoning
    excluded_reason: str | None = None

    # Account-creation request payload (set when review_action =
    # "Pending Account Creation")
    requires_account_creation: bool = False
    new_account_name: str | None = None
    new_account_parent: str | None = None    # resolved with ABBR applied
    new_account_root_type: str | None = None
    new_account_is_group: bool = False

    # Supplier-resolution payload (set when party ledger routed through
    # tier1_supplier). Mutually exclusive with the account fields.
    proposed_supplier: str | None = None       # Supplier ERPNext doc ID
    supplier_match_score: float = 0.0          # 0.0-1.0
    matched_alias_rule: str | None = None      # Supplier Alias Rule name, if any
    requires_supplier_creation: bool = False
    new_supplier_name: str | None = None       # Candidate supplier_name for review


# ---------------------------------------------------------------------------
# Orchestrator
# ---------------------------------------------------------------------------


class Mapper:
    """Resolve Tally ledgers to ERPNext Account proposals.

    Construct once per migration session; call `resolve(ledger)` per ledger
    or `map_all(ledgers)` for a batch. Both are pure functions of the
    (RuleSource, COA, abbr, entity_type) tuple — no hidden state beyond the
    cached rule lists at construction.
    """

    def __init__(
        self,
        rule_source: RuleSource,
        coa: dict[str, CoaAccount],
        abbr: str,
        entity_type: str = "*",
        *,
        supplier_source: SupplierSource | None = None,
        supplier_fuzzy_threshold: float = 85.0,
    ):
        self.rule_source = rule_source
        self.coa = coa
        self.abbr = abbr
        self.entity_type = entity_type
        self.supplier_source = supplier_source
        self.supplier_fuzzy_threshold = supplier_fuzzy_threshold
        self._positive: list[Rule] = rule_source.positive_rules(entity_type)
        self._anti: list[Rule] = rule_source.anti_pattern_rules(entity_type)

    # ---------- public API ----------

    def resolve(self, ledger: Any) -> MappedDecision:
        # 1. P&L exclusion (structural, pre-resolution)
        pnl = pnl_root_type_exclusion(ledger)
        if pnl:
            return self._excluded(ledger, pnl, tier="excluded_pnl")

        # 1b. Party-ledger routing (vendor/creditor → supplier resolution).
        # Only engaged when the Mapper was constructed with a supplier_source.
        # Without a source, party ledgers flow through the account-mapping
        # pathway (and typically land in `unmapped` — matches pre-Work-Item-6
        # behavior).
        if self.supplier_source is not None and is_vendor_party_ledger(ledger):
            return self._resolve_party(ledger)

        # 2. Anti-pattern lookup (does not short-circuit by itself)
        anti = find_matching_anti_pattern(ledger, self._anti)

        # 3. Positive rules (Layers 1 + 2)
        rule_match = find_matching_positive_rule(ledger, self._positive)
        if rule_match:
            rule, tier = rule_match
            return self._resolve_with_rule(ledger, rule, tier, anti)

        # 4. Exact-name fallback (Layer 3)
        exact = resolve_exact_name(ledger, self.coa, self.abbr)
        if exact:
            return self._resolve_with_candidate(ledger, exact, matched_rule=None,
                                                tier="tier1_exact", anti=anti,
                                                confidence=1.0)

        # 5. Anti-pattern matched but no positive target emerged
        if anti:
            return self._anti_pattern_only(ledger, anti)

        # 6. Unmapped
        return MappedDecision(
            tally_name=ledger.name,
            tally_id=ledger.tally_id,
            tally_root_type=ledger.root_type,
            opening_dr=ledger.opening_dr,
            opening_cr=ledger.opening_cr,
            tier="unmapped",
            proposed_account=None,
            review_action="Pending",
            matched_rule=None,
            confidence=0.0,
        )

    def map_all(self, ledgers: Iterable[Any]) -> list[MappedDecision]:
        return [self.resolve(l) for l in ledgers]

    # ---------- internal builders ----------

    def _excluded(
        self, ledger: Any, outcome: ValidatorOutcome, *, tier: str,
    ) -> MappedDecision:
        return MappedDecision(
            tally_name=ledger.name,
            tally_id=ledger.tally_id,
            tally_root_type=ledger.root_type,
            opening_dr=ledger.opening_dr,
            opening_cr=ledger.opening_cr,
            tier=tier,
            proposed_account=None,
            review_action=outcome.review_action,
            matched_rule=None,
            excluded_reason=outcome.excluded_reason,
            confidence=1.0,
        )

    def _resolve_with_rule(
        self, ledger: Any, rule: Rule, tier: str, anti: Rule | None,
    ) -> MappedDecision:
        resolved = (rule.erpnext_account_template or "").replace("{ABBR}", self.abbr)
        if not resolved:
            # Defensive — a positive rule with no template is malformed data.
            return MappedDecision(
                tally_name=ledger.name,
                tally_id=ledger.tally_id,
                tally_root_type=ledger.root_type,
                opening_dr=ledger.opening_dr,
                opening_cr=ledger.opening_cr,
                tier="unmapped",
                proposed_account=None,
                review_action="Pending",
                matched_rule=rule.source_section,
                excluded_reason="Rule has no erpnext_account_template.",
                confidence=0.0,
            )

        # Missing target — creation workflow or unmapped-with-reason
        if resolved not in self.coa:
            return self._missing_target(ledger, rule, resolved, tier, anti)

        return self._resolve_with_candidate(
            ledger, resolved, matched_rule=rule.source_section,
            tier=tier, anti=anti,
            confidence=1.0 if tier == "tier1_rule" else 0.9,
        )

    def _resolve_with_candidate(
        self,
        ledger: Any,
        candidate: str,
        *,
        matched_rule: str | None,
        tier: str,
        anti: Rule | None,
        confidence: float,
    ) -> MappedDecision:
        """Shared post-match plumbing: apply anti-pattern override, then
        group-account refusal, then return the positive decision."""
        if anti and self._template_equals(anti.forbidden_erpnext_template, candidate):
            return self._anti_pattern_block(ledger, anti, blocked_target=candidate)

        gref = group_account_refusal(candidate, self.coa)
        if gref:
            return self._excluded(ledger, gref, tier="group_refused")

        return MappedDecision(
            tally_name=ledger.name,
            tally_id=ledger.tally_id,
            tally_root_type=ledger.root_type,
            opening_dr=ledger.opening_dr,
            opening_cr=ledger.opening_cr,
            tier=tier,
            proposed_account=candidate,
            review_action="Pending",
            matched_rule=matched_rule,
            confidence=confidence,
        )

    def _missing_target(
        self,
        ledger: Any,
        rule: Rule,
        resolved: str,
        tier: str,
        anti: Rule | None,
    ) -> MappedDecision:
        """Rule's resolved target is not in the COA. Either emit a creation
        request (if the rule or the paired anti-pattern requests it), or
        flag as unmapped-target-missing for reviewer attention."""
        creates = rule.creates_erpnext_account
        new_name_tmpl = rule.new_account_name_template
        new_parent_tmpl = rule.new_account_parent
        new_root = rule.new_account_root_type
        new_is_group = rule.new_account_is_group
        anti_blocked = False
        anti_section: str | None = None
        anti_msg: str | None = None

        # Inherit creation from a paired anti-pattern if its suggested
        # alternative matches the positive rule's resolved target.
        if not creates and anti is not None and anti.creates_erpnext_account:
            anti_suggested = (anti.suggested_alternative_template or "").replace(
                "{ABBR}", self.abbr
            )
            if anti_suggested == resolved:
                creates = True
                new_name_tmpl = anti.new_account_name_template
                new_parent_tmpl = anti.new_account_parent
                new_root = anti.new_account_root_type
                new_is_group = anti.new_account_is_group
                anti_blocked = True
                anti_section = anti.source_section
                anti_msg = anti.anti_pattern_reason

        if not creates:
            return MappedDecision(
                tally_name=ledger.name,
                tally_id=ledger.tally_id,
                tally_root_type=ledger.root_type,
                opening_dr=ledger.opening_dr,
                opening_cr=ledger.opening_cr,
                tier="unmapped",
                proposed_account=None,
                review_action="Pending",
                matched_rule=rule.source_section,
                excluded_reason=(
                    f"Rule {rule.source_section} proposes '{resolved}' but the "
                    f"target is not in the COA and the rule has "
                    f"creates_erpnext_account=0."
                ),
                confidence=0.0,
            )

        new_name = (new_name_tmpl or resolved).replace("{ABBR}", self.abbr)
        new_parent = (
            f"{new_parent_tmpl} - {self.abbr}" if new_parent_tmpl else None
        )
        return MappedDecision(
            tally_name=ledger.name,
            tally_id=ledger.tally_id,
            tally_root_type=ledger.root_type,
            opening_dr=ledger.opening_dr,
            opening_cr=ledger.opening_cr,
            tier="pending_account_creation",
            proposed_account=None,
            review_action="Pending Account Creation",
            matched_rule=rule.source_section,
            anti_pattern_blocked=anti_blocked,
            anti_pattern_rule=anti_section,
            anti_pattern_message=anti_msg,
            requires_account_creation=True,
            new_account_name=new_name,
            new_account_parent=new_parent,
            new_account_root_type=new_root,
            new_account_is_group=new_is_group,
            confidence=0.95,
        )

    def _anti_pattern_block(
        self,
        ledger: Any,
        anti: Rule,
        *,
        blocked_target: str,
    ) -> MappedDecision:
        """The mapper found a candidate target that the anti-pattern forbids.
        Route to the suggested alternative — either by proposing the existing
        alternative account, or by emitting a creation request when the
        anti-pattern declares `creates_erpnext_account=1` and the suggested
        alternative is not in the COA."""
        suggested = (anti.suggested_alternative_template or "").replace(
            "{ABBR}", self.abbr
        )

        if suggested and suggested in self.coa:
            # Suggested alternative already exists — propose it with the block flag
            gref = group_account_refusal(suggested, self.coa)
            if gref:
                return self._excluded(ledger, gref, tier="group_refused")
            return MappedDecision(
                tally_name=ledger.name,
                tally_id=ledger.tally_id,
                tally_root_type=ledger.root_type,
                opening_dr=ledger.opening_dr,
                opening_cr=ledger.opening_cr,
                tier="anti_pattern_blocked",
                proposed_account=suggested,
                review_action="Pending",
                matched_rule=None,
                anti_pattern_blocked=True,
                anti_pattern_rule=anti.source_section,
                anti_pattern_message=anti.anti_pattern_reason,
                excluded_reason=f"Refused forbidden target '{blocked_target}'; "
                                f"steered to '{suggested}'.",
                confidence=0.9,
            )

        if anti.creates_erpnext_account:
            new_name = (
                (anti.new_account_name_template or suggested or "")
                .replace("{ABBR}", self.abbr)
            )
            new_parent = (
                f"{anti.new_account_parent} - {self.abbr}"
                if anti.new_account_parent else None
            )
            return MappedDecision(
                tally_name=ledger.name,
                tally_id=ledger.tally_id,
                tally_root_type=ledger.root_type,
                opening_dr=ledger.opening_dr,
                opening_cr=ledger.opening_cr,
                tier="pending_account_creation",
                proposed_account=None,
                review_action="Pending Account Creation",
                matched_rule=None,
                anti_pattern_blocked=True,
                anti_pattern_rule=anti.source_section,
                anti_pattern_message=anti.anti_pattern_reason,
                excluded_reason=(
                    f"Refused forbidden target '{blocked_target}'. Suggested "
                    f"alternative '{new_name}' does not exist in COA; "
                    f"creation requested."
                ),
                requires_account_creation=True,
                new_account_name=new_name,
                new_account_parent=new_parent,
                new_account_root_type=anti.new_account_root_type,
                new_account_is_group=anti.new_account_is_group,
                confidence=0.95,
            )

        # No suggested alternative resolvable — pure block, reviewer decides
        return MappedDecision(
            tally_name=ledger.name,
            tally_id=ledger.tally_id,
            tally_root_type=ledger.root_type,
            opening_dr=ledger.opening_dr,
            opening_cr=ledger.opening_cr,
            tier="anti_pattern_blocked",
            proposed_account=None,
            review_action="Pending",
            matched_rule=None,
            anti_pattern_blocked=True,
            anti_pattern_rule=anti.source_section,
            anti_pattern_message=anti.anti_pattern_reason,
            excluded_reason=(
                f"Anti-pattern {anti.source_section} refused '{blocked_target}'; "
                f"no resolvable alternative."
            ),
            confidence=0.0,
        )

    def _resolve_party(self, ledger: Any) -> MappedDecision:
        """Party-ledger (Sundry Creditors descendant) → Supplier."""
        assert self.supplier_source is not None  # caller guards
        tier, supplier, score, alias_rule = resolve_supplier(
            ledger,
            self.supplier_source,
            fuzzy_threshold=self.supplier_fuzzy_threshold,
        )

        base = dict(
            tally_name=ledger.name,
            tally_id=ledger.tally_id,
            tally_root_type=ledger.root_type,
            opening_dr=ledger.opening_dr,
            opening_cr=ledger.opening_cr,
        )

        if tier == "pending_supplier_creation":
            # No match at any layer — emit a Supplier Creation Request stub.
            # new_supplier_name carries the cleaned Tally ledger name as the
            # reviewer's suggested starting point; reviewer edits as needed.
            from rgi_migration.mapper.tier1_supplier import _clean
            return MappedDecision(
                **base,
                tier="pending_supplier_creation",
                proposed_account=None,
                review_action="Pending Supplier Creation",
                matched_rule=None,
                confidence=0.0,
                requires_supplier_creation=True,
                new_supplier_name=_clean(ledger.name),
            )

        # Matched (exact, alias, or fuzzy)
        return MappedDecision(
            **base,
            tier=tier,
            proposed_account=None,
            review_action="Pending",
            matched_rule=None,
            confidence=score,
            proposed_supplier=supplier.name if supplier else None,
            supplier_match_score=score,
            matched_alias_rule=alias_rule,
        )

    def _anti_pattern_only(self, ledger: Any, anti: Rule) -> MappedDecision:
        """Anti-pattern matched but no positive rule produced a target and
        exact-name fallback did not hit. Treat the forbidden template as the
        'candidate we would have picked' and funnel through the block path."""
        forbidden = (anti.forbidden_erpnext_template or "").replace(
            "{ABBR}", self.abbr
        )
        return self._anti_pattern_block(ledger, anti, blocked_target=forbidden or "<unspecified>")

    @staticmethod
    def _template_equals(template: str | None, resolved: str) -> bool:
        """Template-vs-resolved comparison that treats `{ABBR}` as a wildcard."""
        if not template:
            return False
        if "{ABBR}" not in template:
            return template == resolved
        before, after = template.split("{ABBR}", 1)
        return resolved.startswith(before) and resolved.endswith(after)


# ---------------------------------------------------------------------------
# Reporting
# ---------------------------------------------------------------------------


def summarize(decisions: list[MappedDecision]) -> dict[str, Any]:
    """Compact hit-rate summary suitable for CLI stdout or Session snapshot."""
    by_tier: Counter = Counter(d.tier for d in decisions)
    by_action: Counter = Counter(d.review_action for d in decisions)
    total = len(decisions)
    unmapped = by_tier.get("unmapped", 0)
    anti = sum(1 for d in decisions if d.anti_pattern_blocked)
    account_creations = sum(1 for d in decisions if d.requires_account_creation)
    supplier_creations = sum(1 for d in decisions if d.requires_supplier_creation)
    party_ledgers = sum(
        1 for d in decisions
        if d.tier in (
            "tier1_supplier_exact",
            "tier1_supplier_alias",
            "tier1_supplier_fuzzy",
            "pending_supplier_creation",
        )
    )
    return {
        "total": total,
        "by_tier": dict(by_tier),
        "by_review_action": dict(by_action),
        "anti_pattern_blocked_count": anti,
        "account_creation_requests": account_creations,
        "supplier_creation_requests": supplier_creations,
        "party_ledgers_routed": party_ledgers,
        "hit_rate": round((total - unmapped) / max(1, total), 3),
    }


# ---------------------------------------------------------------------------
# COA loader
# ---------------------------------------------------------------------------


def _truthy(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    return str(value).strip() in ("1", "True", "true", "yes", "Yes")


def _pick(row: dict[str, Any], *keys: str) -> Any:
    """First non-empty value among the given keys (case-sensitive match on
    whatever the source file supplied)."""
    for k in keys:
        if k in row:
            v = row[k]
            if v is None:
                continue
            if isinstance(v, str) and not v.strip():
                continue
            return v
    return ""


def load_coa(path: Path) -> dict[str, CoaAccount]:
    """Load the target company's ERPNext COA from CSV or XLSX.

    Accepts two header schemas so both the hand-built stub and a raw
    ERPNext export can be used directly (no munging step):

    Internal / stub schema:
        account_name, parent_account, root_type, is_group, company_abbr

    ERPNext native export schema (Desk → Accounts → Export):
        ID, Account Name, Company, Parent Account, Is Group, Root Type,
        Account Category, Account Type

    In the native schema, the fully-qualified account name (what the mapper
    keys on) lives in the `ID` column — `Account Name` is the base name
    without the `- {ABBR}` suffix. Extra columns are ignored.
    """
    path = Path(path)
    suffix = path.suffix.lower()
    rows: list[dict[str, Any]] = []

    if suffix == ".csv":
        with path.open(encoding="utf-8") as f:
            rows = list(csv.DictReader(f))
    elif suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook  # type: ignore[import]
        except ImportError as e:   # pragma: no cover
            raise RuntimeError("openpyxl required to load an XLSX COA") from e
        wb = load_workbook(str(path), data_only=True)
        ws = wb.active
        header = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
        for r in ws.iter_rows(min_row=2, values_only=True):
            row = {
                header[i]: ("" if v is None else v)
                for i, v in enumerate(r) if i < len(header)
            }
            rows.append(row)
    else:
        raise ValueError(f"Unsupported COA file extension: {suffix}")

    out: dict[str, CoaAccount] = {}
    for r in rows:
        # Fully-qualified account name: `account_name` (internal) or `ID` (ERPNext export)
        name = str(_pick(r, "account_name", "ID", "id")).strip()
        if not name:
            continue
        parent = str(_pick(r, "parent_account", "Parent Account")).strip() or None
        root = str(_pick(r, "root_type", "Root Type")).strip()
        is_group_raw = _pick(r, "is_group", "Is Group")
        # ERPNext export uses "Company" (full name); stub uses "company_abbr" (short).
        # Either is fine as metadata; mapper takes the abbr from the --abbr CLI flag.
        company_abbr = str(_pick(r, "company_abbr", "Company")).strip()
        out[name] = CoaAccount(
            name=name,
            parent_account=parent,
            root_type=root,
            is_group=_truthy(is_group_raw),
            company_abbr=company_abbr,
        )
    return out


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def _cli(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(
        prog="python -m rgi_migration.mapper.mapper",
        description=(
            "Resolve Tally ledgers against the target company's ERPNext COA "
            "using the Tier-1 rule library."
        ),
    )
    p.add_argument("--coa", required=True, metavar="PATH",
                   help="ERPNext COA (CSV or XLSX) with columns "
                        "account_name, parent_account, root_type, is_group, "
                        "company_abbr.")
    p.add_argument("--rules", default="docs/seed_plan.json", metavar="PATH",
                   help="Path to seed_plan.json (default: docs/seed_plan.json).")
    p.add_argument("--abbr", required=True, metavar="ABBR",
                   help="Target company abbreviation (e.g. CACSPU).")
    p.add_argument("--entity-type", default="*", metavar="TYPE",
                   help="Entity-type filter (college | hostel | society | "
                        "university | hospital | * ). Default: *.")
    p.add_argument("--suppliers", metavar="PATH", default=None,
                   help="Optional: Supplier master CSV (e.g. "
                        "rgi_migration/tests/fixtures/jewonline_suppliers_real.csv). "
                        "When provided, party ledgers (Sundry Creditors descendants) "
                        "are routed through Tier-1 supplier resolution.")
    p.add_argument("--supplier-fuzzy-threshold", type=float, default=85.0,
                   help="Fuzzy-match threshold for the Layer-3 general supplier "
                        "fallback (default 85.0).")
    src = p.add_mutually_exclusive_group(required=True)
    src.add_argument("--xml", metavar="PATH", help="Tally All Masters XML to parse.")
    src.add_argument("--excel", metavar="PATH", help="Tally opening-TB Excel to parse.")
    p.add_argument("--verbose", action="store_true")
    args = p.parse_args(argv)

    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(levelname)s %(name)s: %(message)s",
    )

    coa = load_coa(Path(args.coa))
    LOG.info("Loaded %d COA entries from %s", len(coa), args.coa)

    rules = JsonFileRuleSource(Path(args.rules))
    LOG.info(
        "Loaded %d positive + %d anti-pattern rules (status=confirmed, "
        "entity_type=%s)",
        len(rules.positive_rules(args.entity_type)),
        len(rules.anti_pattern_rules(args.entity_type)),
        args.entity_type,
    )

    if args.xml:
        from rgi_migration.parsers.tally_xml_parser import parse_xml  # type: ignore[import]
        tb = parse_xml(args.xml)
    else:
        from rgi_migration.parsers.tally_excel_parser import parse_excel  # type: ignore[import]
        tb = parse_excel(args.excel)
    LOG.info(
        "Parsed %d main ledgers from Tally source (+%d student, +%d system)",
        len(tb.ledgers), len(tb.student_ledgers), len(tb.system_ledgers),
    )

    supplier_source = None
    if args.suppliers:
        from rgi_migration.mapper.supplier_source import CsvFileSupplierSource
        supplier_source = CsvFileSupplierSource(args.suppliers)
        LOG.info(
            "Loaded %d suppliers from %s (fuzzy threshold %.1f)",
            len(supplier_source.get_all_suppliers()),
            args.suppliers, args.supplier_fuzzy_threshold,
        )

    mapper = Mapper(
        rules, coa, abbr=args.abbr, entity_type=args.entity_type,
        supplier_source=supplier_source,
        supplier_fuzzy_threshold=args.supplier_fuzzy_threshold,
    )
    decisions = mapper.map_all(tb.ledgers)
    summary = summarize(decisions)
    print(json.dumps(summary, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(_cli())
